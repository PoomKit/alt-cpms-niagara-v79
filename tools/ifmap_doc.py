# -*- coding: utf-8 -*-
"""Field interface -> logic block map document (Thai, A4 landscape).  usage: ifmap_doc.py <ifmap outdir> <out.docx>"""
import sys, os, json, collections
import openpyxl
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

D, OUT = sys.argv[1:3]
FONT = os.environ.get('MANUAL_FONT', 'Leelawadee UI'); SC = float(os.environ.get('MANUAL_SCALE', '1'))
doc = Document(); sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Cm(1.4))
W = 29.7 - 2.8


def style(st, size, bold=None, color=None):
    st.font.name = FONT; st.font.size = Pt(size * SC)
    if bold is not None: st.font.bold = bold
    if color: st.font.color.rgb = RGBColor.from_string(color)
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(k), FONT)
    e = OxmlElement('w:szCs'); e.set(qn('w:val'), str(int(size * SC * 2))); rpr.append(e)
    if bold: rpr.append(OxmlElement('w:bCs'))


style(doc.styles['Normal'], 10)
for l, s_, c in ((1, 16, '1F4E79'), (2, 13, '1F4E79'), (3, 11, '2E5E8C')): style(doc.styles[f'Heading {l}'], s_, True, c)


def run(p, text, size=None, bold=False, color=None):
    r = p.add_run(text); r.bold = bold
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size * SC); rpr = r._r.get_or_add_rPr(); e = OxmlElement('w:szCs'); e.set(qn('w:val'), str(int(size * SC * 2))); rpr.append(e)
    return r


def P(text, size=None, bold=False, color=None):
    p = doc.add_paragraph()
    for i, part in enumerate(text.split('**')): run(p, part, size, bold or i % 2 == 1, color)
    return p


def BL(items, style_='List Bullet'):
    for t in items:
        p = doc.add_paragraph(style=style_)
        for i, part in enumerate(t.split('**')): run(p, part, None, i % 2 == 1)


def shade(c, h):
    tcPr = c._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), h); tcPr.append(s)


def T(head, rows, widths, size=8, fills=None):
    t = doc.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for c, h, w in zip(t.rows[0].cells, head, widths):
        c.width = Cm(w); run(c.paragraphs[0], h, size, True, 'FFFFFF'); shade(c, '1F4E79')
    for k, r in enumerate(rows):
        for j, (c, v, w) in enumerate(zip(t.add_row().cells, r, widths)):
            c.width = Cm(w); run(c.paragraphs[0], str(v), size, j == 0 and fills is not None)
            if fills and fills.get(r[0]) and j == 0: shade(c, fills[r[0]])
            elif k % 2: shade(c, 'F2F5F8')
    doc.add_paragraph()


M = json.load(open(os.path.join(D, 'map.json'), encoding='utf-8'))
wb = openpyxl.load_workbook(os.path.join(D, 'V7_7_INTERFACE_LOGIC_MAP.xlsx'))
pts = list(wb['Points'].values)[1:]
n_in = sum(1 for r in pts if r[1].startswith('IN')); n_out = len(pts) - n_in

doc.add_heading('Field Interface → Logic Block Map', 0)
P('ALT_CP_Ref_V7_7 (r7)  ·  REFERENCE SIMULATION – NOT PROJECT APPROVED', 10, True, 'B71C1C')
P(f'เอกสารนี้บอกว่า field point แต่ละจุด ({len(pts)} จุด: input {n_in} / output {n_out}) ต้องต่อเข้าที่ไหน และไปเข้า logic block ตัวไหน '
  'ใช้ตอนนำ logic ไปใช้กับอุปกรณ์จริง ตรวจ P2P และแก้ปัญหา ไฟล์ Excel `V7_7_INTERFACE_LOGIC_MAP.xlsx` มีข้อมูลเดียวกันแบบ filter ได้')

doc.add_heading('1. หลักการ: หน้างานต่อ link แค่จุดเดียวต่อ point', 1)
BL(['**INPUT (AI / DI):** อุปกรณ์จริง → driver proxy point → **IO_<point>** → IOS_<point> (เลือก FIELD / SIM) → <point>.in10 → logic block ที่อ่านค่า',
    '**OUTPUT (DO / AO):** logic block → <point>.in16 → IOS_<point> (OutEnable) → **IO_<point>**.in10 → IO_<point>.out → driver proxy point → อุปกรณ์จริง',
    '**สิ่งที่หน้างานทำ:** สร้าง proxy point ตาม address แล้ว link กับ IO_<point> เท่านั้น (input: proxy.out → IO_.in16 · output: IO_.out → proxy.in10)',
    '**ห้าม** link proxy เข้า logic block หรือ logic point โดยตรง เพราะจะข้ามตัวเลือก FIELD / SIM, P2P isolate และ fallback',
    '**AI (heartbeat + setpoint 3 ค่า) ไม่อยู่ในเอกสารนี้:** รับจาก cloud ตาม tag melco:pointNameString ไม่ใช้สาย BACnet ดูเอกสาร AI Cloud Interface – Tag Map',
    'IO_ ทุกจุดอยู่ในโฟลเดอร์ **FieldIO/<group>** (ChillerPlant/FieldIO และ AirSide/FieldIO) แยกตามเครื่อง ตรงกับหัวข้อที่ 3'])
doc.add_picture(os.path.join(D, 'img', 'IF_00_example.png'), width=Cm(W))
P('อ่านรูป: ช่อง 1 = address ที่ต้องสร้าง proxy · ช่อง 2 = IO_ point ที่ต้อง link (ตัวแดงบอก slot) · ช่อง 3 = ตัวเลือก FIELD/SIM · ช่อง 4 = logic point · '
  'ช่อง 5 = logic block ที่ใช้ค่า (input) หรือที่สั่งค่า (output) พร้อมชนิด block, slot ที่เข้า และโฟลเดอร์', 9)
T(['สีใน ช่อง 5', 'ความหมาย'], [['Control', 'ควบคุมจริง (sequence, PID, staging, interlock ของการเดินเครื่อง)'],
                                 ['Alarm / safety', 'alarm และ trip / interlock ความปลอดภัย'], ['Smart O&M', 'PM, FDD (แจ้งเตือนล่วงหน้า), Health, demand limit'],
                                 ['Commissioning test', 'ใช้ตอน functional test เท่านั้น'], ['Graphics (display)', 'คำนวณเพื่อแสดงบน Px'],
                                 ['Simulation only', 'model จำลอง ไม่มีผลเมื่อเป็น FIELD']], (5, W - 5), 9,
  {'Control': 'BBDEFB', 'Alarm / safety': 'FFCDD2', 'Smart O&M': 'E1BEE7', 'Commissioning test': 'FFE0B2', 'Graphics (display)': 'B2DFDB', 'Simulation only': 'E0E0E0'})

doc.add_heading('2. ขั้นตอนเปิดใช้ field point', 1)
BL(['เพิ่ม driver / device ตาม Site Deployment Design (BACnet/IP, Modbus RTU, MS/TP, IO module)',
    'discover / สร้าง proxy point ตาม address ในช่อง 1 (ตั้งชื่อ proxy เหมือน point เพื่อค้นง่าย)',
    'link proxy กับ IO_<point> ตามช่อง 2 ทีละกลุ่ม (ใช้ Excel sheet Points filter ตาม Group)',
    'ตั้ง IO_<group>_Protocol ให้ตรงกับที่ใช้จริง และตั้ง **IO_DriversInstalled = true**',
    'ทีละกลุ่ม: IO_<group>_UseField = true → ทำ P2P (หน้า Commissioning) → ดูว่า logic block ในช่อง 5 ได้ค่าจริง',
    'ครบทุกกลุ่มแล้วจึงตั้ง IO_AllField = true แล้วรัน functional test FT01–FT10'], 'List Number')
nl = M['nologic']
P(f'**Point ที่ไม่มี logic ใช้ ({len(nl)} จุด):** เป็นค่าแสดงผล / history / energy เท่านั้น (เช่น Amps, Volts, VFD DC bus, Hz feedback) '
  'ต่อหรือไม่ต่อก็ไม่กระทบการควบคุม ควรต่อเพื่อ report และ FDD ในอนาคต', 9)

doc.add_page_break()
doc.add_heading('3. แผนผังต่อกลุ่มอุปกรณ์', 1)
last = None
for g, fn in M['imgs']:
    if g != last:
        doc.add_heading(f'{g}  ({M["groups"][g]} points)', 2); last = g
    from PIL import Image
    w, h = Image.open(fn).size
    if h / w * W > 16.0: doc.add_picture(fn, height=Cm(16.0))
    else: doc.add_picture(fn, width=Cm(W))
    doc.add_page_break()

doc.add_heading('4. Logic function → field point ที่ต้องใช้', 1)
P('ใช้ดูกลับด้าน: ถ้าจะทดสอบ / เปิดใช้ logic ในโฟลเดอร์ไหน ต้อง interface point ไหนก่อน (เฉพาะ Control, Alarm / safety, Smart O&M · รายละเอียดราย block อยู่ใน Excel sheet Logic→Points)', 9)
fs = [r for r in list(wb['Function summary'].values)[1:] if r[0] in ('Control', 'Alarm / safety', 'Smart O&M') and (r[3] or r[5])]
T(['Area', 'Logic folder', 'Input ที่ต้องต่อ', 'Output ที่สั่ง'], [[r[0], r[1], r[4] or '-', r[6] or '-'] for r in fs], (2.6, 5.4, 11.5, W - 19.5), 7,
  {'Control': 'BBDEFB', 'Alarm / safety': 'FFCDD2', 'Smart O&M': 'E1BEE7'})
doc.save(OUT); print('saved', OUT)
