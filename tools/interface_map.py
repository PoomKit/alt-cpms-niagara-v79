# -*- coding: utf-8 -*-
"""Field interface -> logic block map for ALT_CP_Ref_V7_7.
For every real field point: field address -> driver proxy -> IO_ landing point -> IOS_ selector -> logic point -> logic blocks
(inputs: blocks that READ the point; outputs: blocks that DRIVE the command).  Reverse view: logic block -> field points it needs.
usage: interface_map.py <config.bog> <regdir> <outdir>      -> V7_7_INTERFACE_LOGIC_MAP.xlsx, img/*.png, map.json"""
import sys, os, csv, json, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ws_manual as WS
from PIL import Image, ImageDraw, ImageFont

BOG, REGDIR, OUT = sys.argv[1:4]
os.makedirs(os.path.join(OUT, 'img'), exist_ok=True)
comp, fol, links = WS.load(BOG)
WS.load_kinds(REGDIR)
rev = collections.defaultdict(list)
for t, lk in links.items():
    for slot, src, _ in lk:
        if src: rev[src].append((t, slot))
reg = list(csv.DictReader(open(os.path.join(REGDIR, 'V7_7_FIELD_IO_REGISTER.csv'), encoding='utf-8-sig')))
note = {r['Point']: r['Note'] for r in csv.DictReader(open(os.path.join(REGDIR, 'V7_POINT_LIST.csv'), encoding='utf-8-sig'))}
SLOTP = {'in1': 'in1 emergency', 'in8': 'in8 manual/override', 'in9': 'in9 test (CX)', 'in10': 'in10 field/sim', 'in16': 'in16 logic'}
AREAS = ['Control', 'Alarm / safety', 'Smart O&M', 'Commissioning test', 'Graphics (display)', 'Simulation only']


def area(n):
    f = fol.get(n, '')
    if n.startswith('CX_') or f.startswith('Commissioning'): return 'Commissioning test'
    if f.startswith('Smart') or n.startswith(('PM_', 'FDD_', 'Smart_', 'DL_')): return 'Smart O&M'
    if n.startswith('Sim_') or 'Simulation' in f or 'Model' in f: return 'Simulation only'
    if n.startswith('Gfx_') or f.startswith('Graphics'): return 'Graphics (display)'
    if 'Alarm' in f or n.endswith(('Alarm', 'Alm', 'Trip', 'TripLatch')) or 'Safety' in f: return 'Alarm / safety'
    return 'Control'


def func(n):
    el = comp.get(n)
    if el is None: return ''
    lk = links.get(n, [])
    try: return WS.function_line(n, el, lk, WS.category(n, el, lk))
    except Exception: return el.get('t').split(':')[1]


def btype(n):
    el = comp.get(n)
    return el.get('t').split(':')[1] if el is not None else '?'


def addr(r):
    p = r['Protocol_Name']
    if p.startswith('BACnet'): return f"{p} dev {r['BACnet_Device']} {r['BACnet_Object']}"
    if p.startswith('Modbus'): return f"{p} unit {r['Modbus_Unit']} {r['Modbus_Register']}"
    return f"{p} {r['IO_Module']} {r['IO_Channel']}"


def block(n, slot):
    return dict(name=n, slot=slot, area=area(n), type=btype(n), folder='/' + fol.get(n, ''), func=func(n))


ORDER = {a: i for i, a in enumerate(AREAS)}
rows = []
for r in reg:
    pt, d = r['Point'], r['Direction']
    if d == 'in':
        main = [block(t, s) for t, s in rev.get(pt, []) if not t.startswith('IOS_')]
        also = []
        site = f"สร้าง proxy point ({addr(r)}) แล้ว link  proxy.out → {r['IO_Point']}.in16"
    else:
        main = [block(src, SLOTP.get(s, s)) for s, src, _ in links.get(pt, []) if src and not src.startswith('IOS_')]
        also = [block(t, s) for t, s in rev.get(pt, []) if not t.startswith('IOS_')]
        site = f"สร้าง proxy point แบบ writable ({addr(r)}) แล้ว link  {r['IO_Point']}.out → proxy.in10"
    main.sort(key=lambda b: (ORDER[b['area']], b['folder'], b['name']))
    also.sort(key=lambda b: (ORDER[b['area']], b['folder'], b['name']))
    rows.append(dict(r=r, main=main, also=also, site=site, desc=note.get(pt, '')))

# ------------------------------------------------------------------ XLSX
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
wb = openpyxl.Workbook()
HDR = PatternFill('solid', fgColor='1F4E79'); WH = Font(bold=True, color='FFFFFF'); B = Font(bold=True)
FILL = {'in': PatternFill('solid', fgColor='E3F2FD'), 'out': PatternFill('solid', fgColor='FFF3E0')}
STEPF = [PatternFill('solid', fgColor=c) for c in ('ECEFF1', 'E8F5E9', 'F3E5F5', 'FFFDE7', 'E3F2FD')]
thin = Side(style='thin', color='C9D1D9'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WR = Alignment(wrap_text=True, vertical='top')


def sheet(ws, head, data, widths, fills=None):
    ws.append(head)
    for c in ws[1]: c.font = WH; c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(data):
        ws.append(row)
        for j, c in enumerate(ws[ws.max_row]):
            c.alignment = WR; c.border = BOX
            if fills:
                f = fills(i, j)
                if f: c.fill = f
    for i, w in enumerate(widths, 1): ws.column_dimensions[CL(i)].width = w
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 42


ws = wb.active; ws.title = 'วิธีต่อ (How to)'
howto = [
    ['FIELD INTERFACE → LOGIC BLOCK MAP  ·  ALT_CP_Ref_V7_7  ·  REFERENCE – NOT PROJECT APPROVED'],
    [''],
    ['เส้นทางสัญญาณ INPUT (AI / DI) : อุปกรณ์จริง → Driver proxy point → IO_<point> → IOS_<point> (FIELD/SIM) → <point>.in10 → logic block ที่อ่านค่า'],
    ['เส้นทางสัญญาณ OUTPUT (DO / AO): logic block → <point>.in16 → IOS_<point> (OutEnable) → IO_<point>.in10 → IO_<point>.out → Driver proxy point → อุปกรณ์จริง'],
    [''],
    ['หน้างานต่อ link แค่จุดเดียวต่อ point: ระหว่าง proxy point กับ IO_<point> (คอลัมน์ "หน้างานต้องทำ") ห้ามต่อ proxy เข้า logic block โดยตรง'],
    ['IOS_ / logic point / logic block ไม่ต้องแก้ ต่อไว้ครบแล้ว ตารางนี้บอกว่า IO_ แต่ละจุดไปเข้า logic block ไหน เพื่อใช้ตรวจ / ทดสอบ / แก้ปัญหา'],
    [''],
    ['ลำดับเปิดใช้งานจริง'],
    ['1. เพิ่ม driver (BacnetNetwork / ModbusAsyncNetwork / IO module) และ device ตาม Design document'],
    ['2. discover / สร้าง proxy point ตามคอลัมน์ Field address'],
    ['3. link proxy ↔ IO_<point> ตามคอลัมน์ "หน้างานต้องทำ" (ใช้ sheet "Points" filter ตาม Group)'],
    ['4. ตั้ง IO_<group>_Protocol ให้ตรงกับที่ใช้จริง'],
    ['5. ตั้ง IO_DriversInstalled = true (ถ้ายังไม่ตั้ง ระบบยังเป็น SIMULATION แม้เลือก FIELD)'],
    ['6. ทีละกลุ่ม: IO_<group>_UseField = true → ทำ P2P (Commissioning) → ดู logic block ในคอลัมน์ Logic blocks ว่าได้ค่าจริง'],
    ['7. ครบทุกกลุ่มแล้วจึงใช้ IO_AllField = true'],
    [''],
    ['Sheets: Points = ทุก field point (1 แถว = 1 point)  ·  Logic→Points = logic block แต่ละตัวต้องใช้ point ไหน  ·  Function summary = ต่อโฟลเดอร์ logic  ·  Groups = จำนวนต่อกลุ่ม'],
    ['Logic area: Control = ควบคุมจริง · Alarm / safety = alarm และ interlock · Smart O&M = PM / FDD / Health · Commissioning test = ใช้ตอนทดสอบ · Simulation only = ไม่มีผลหน้างาน'],
]
for h in howto: ws.append(h)
ws['A1'].font = Font(bold=True, size=13, color='1F4E79'); ws['A9'].font = B
for c in ('A3', 'A4'): ws[c].font = Font(bold=True, color='0D47A1' if c == 'A3' else 'E65100')
ws['A6'].font = Font(bold=True, color='B71C1C'); ws.column_dimensions['A'].width = 160


def fmt_blocks(bs, n=99):
    out = [f"{b['name']}  [{b['area']}]  {b['slot']}  ·  {b['type']}: {b['func']}  ·  {b['folder']}" for b in bs[:n]]
    return '\n'.join(out) if out else '— (แสดงผล / history เท่านั้น ไม่มี logic ใช้)'


head = ['Group', 'Dir', 'Type', 'Field point (logic point)', 'Description', '1 Field address', 'Signal',
        '2 หน้างานต้องทำ (link นี้จุดเดียว)', '2 IO_ landing point', '3 Selector (IOS_)', '4 Logic point path',
        '5 Logic blocks (input: อ่านค่า · output: สั่งค่า)', 'Control blocks', 'Alarm', 'Smart', 'Test', 'Also uses this command (output)']
data = []
for x in rows:
    r = x['r']; c = collections.Counter(b['area'] for b in x['main'])
    data.append([r['Group'], 'IN →' if r['Direction'] == 'in' else '← OUT', r['Type'] + (f" ({r['Unit']})" if r['Unit'] else ''),
                 r['Point'], x['desc'], addr(r), r['Signal'], x['site'], r['IO_Path'], 'IOS_' + r['Point'], r['Point_Path'],
                 fmt_blocks(x['main']), c['Control'], c['Alarm / safety'], c['Smart O&M'], c['Commissioning test'], fmt_blocks(x['also'], 8) if x['also'] else ''])
sheet(wb.create_sheet('Points'), head, data, (9, 7, 9, 30, 34, 30, 15, 46, 44, 26, 44, 95, 8, 7, 7, 7, 60),
      lambda i, j: FILL[rows[i]['r']['Direction']] if j < 4 else (STEPF[0] if j in (5, 6) else STEPF[1] if j == 7 else STEPF[2] if j in (8, 9) else STEPF[3] if j == 10 else None))

lp = collections.defaultdict(lambda: {'in': [], 'out': [], 'b': None})
for x in rows:
    d = x['r']['Direction']
    for b in x['main']:
        e = lp[b['name']]; e['b'] = b; e[d].append(f"{x['r']['Point']}  ({x['r']['IO_Point']} · {x['r']['Type']})")
data2 = sorted(([v['b']['area'], v['b']['folder'], k, v['b']['type'], v['b']['func'], '\n'.join(v['in']), '\n'.join(v['out']), len(v['in']) + len(v['out'])]
                for k, v in lp.items()), key=lambda r: (ORDER[r[0]], r[1], r[2]))
sheet(wb.create_sheet('Logic→Points'), ['Area', 'Logic folder (function)', 'Logic block', 'Block type', 'ทำหน้าที่',
                                        'ต้องใช้ field INPUT (ต่อที่ IO_)', 'สั่ง field OUTPUT (ออกที่ IO_)', 'จำนวน'], data2, (16, 40, 34, 16, 40, 52, 52, 8))
fs = collections.defaultdict(lambda: [set(), set(), 0])
for r in data2:
    e = fs[(r[0], r[1])]; e[2] += 1
    for s in r[5].split('\n'):
        if s: e[0].add(s.split('  (')[0])
    for s in r[6].split('\n'):
        if s: e[1].add(s.split('  (')[0])
data3 = [[a, f, n, len(i), ', '.join(sorted(i)), len(o), ', '.join(sorted(o))] for (a, f), (i, o, n) in sorted(fs.items(), key=lambda k: (ORDER[k[0][0]], k[0][1]))]
sheet(wb.create_sheet('Function summary'), ['Area', 'Logic folder (function)', 'Logic blocks', 'Input points', 'Input points list', 'Output points', 'Output points list'],
      data3, (16, 42, 10, 10, 90, 10, 60))
gs = collections.Counter((x['r']['Group'], x['r']['Protocol_Name'], x['r']['Direction'], x['r']['Type']) for x in rows)
sheet(wb.create_sheet('Groups'), ['Group', 'Protocol', 'Dir', 'Type', 'Points'], [list(k) + [v] for k, v in sorted(gs.items())], (14, 18, 8, 8, 8))
xl = os.path.join(OUT, 'V7_7_INTERFACE_LOGIC_MAP.xlsx'); wb.save(xl)

# ------------------------------------------------------------------ diagrams (one row per point)
F = ImageFont.truetype
DJC, DJCB = WS.DJC, WS.DJCB
fName, fSm, fTh, fThB, fHd = F(DJCB, 15), F(DJC, 13), F(WS.SAR, 15), F(WS.SARB, 17), F(WS.SARB, 22)
COLS = [('1  Field device / address', 360, '#ECEFF1', '#607D8B'), ('2  IO_ landing point (FieldIO)', 330, '#E8F5E9', '#2E7D32'),
        ('3  IOS_', 120, '#F3E5F5', '#7B1FA2'), ('4  Logic point', 330, '#FFFDE7', '#F9A825'),
        ('5  Logic blocks ที่ใช้ point นี้', 700, '#FFFFFF', '#1565C0')]
GAP, PAD, LH = 46, 24, 21
ACOL = {'Control': '#1565C0', 'Alarm / safety': '#C62828', 'Smart O&M': '#6A1B9A', 'Commissioning test': '#EF6C00', 'Graphics (display)': '#00897B', 'Simulation only': '#9E9E9E'}
WIDTH = PAD * 2 + sum(c[1] for c in COLS) + GAP * (len(COLS) - 1)


def fit(d, s, f, w):
    if d.textlength(s, font=f) <= w: return s
    while s and d.textlength(s + '…', font=f) > w: s = s[:-1]
    return s + '…'


def arrow(d, x1, x2, y, col, right):
    d.line((x1, y, x2, y), fill=col, width=4)
    tip, base = (x2, x2 - 14) if right else (x1, x1 + 14)
    d.polygon([(tip, y), (base, y - 8), (base, y + 8)], fill=col)


def row_h(x):
    return max(3, min(len(x['main']), 4) + (1 if len(x['main']) > 4 else 0)) * LH + 18


def draw_rows(title, sub, rs, path):
    H = 150 + sum(row_h(x) + 10 for x in rs) + 60
    im = Image.new('RGB', (WIDTH, H), 'white'); d = ImageDraw.Draw(im)
    d.text((PAD, 14), title, font=fHd, fill='#1F4E79'); d.text((PAD, 50), sub, font=fTh, fill='#37474F')
    xs, x = [], PAD
    for name, w, bg, fg in COLS:
        xs.append(x); d.rectangle((x, 86, x + w, 128), fill=fg); d.text((x + 8, 92), name, font=fThB, fill='white'); x += w + GAP
    y = 140
    for k, r in enumerate(rs):
        R = r['r']; h = row_h(r); inp = R['Direction'] == 'in'; ac = '#1565C0' if inp else '#E65100'
        if k % 2: d.rectangle((0, y - 5, WIDTH, y + h + 5), fill='#FAFAFA')
        for i, (name, w, bg, fg) in enumerate(COLS[:4]):
            d.rounded_rectangle((xs[i], y, xs[i] + w, y + h), 8, fill=bg, outline=fg, width=2)
        cx = lambda i: xs[i] + 10
        tag = ('IN ' if inp else 'OUT ') + R['Type']
        d.rounded_rectangle((cx(0), y + 8, cx(0) + 70, y + 30), 5, fill=ac); d.text((cx(0) + 6, y + 9), tag, font=F(DJCB, 13), fill='white')
        d.text((cx(0) + 80, y + 9), fit(d, R['Protocol_Name'], fName, 260), font=fName, fill='#263238')
        a = addr(R).replace(R['Protocol_Name'], '').strip()
        d.text((cx(0), y + 36), fit(d, a, fSm, 340), font=fSm, fill='#37474F')
        d.text((cx(0), y + 36 + LH), fit(d, R['Signal'], fSm, 340), font=fSm, fill='#78909C')
        d.text((cx(1), y + 9), fit(d, R['IO_Point'], fName, 310), font=fName, fill='#1B5E20')
        d.text((cx(1), y + 34), fit(d, R['IO_Path'].rsplit('/', 1)[0], fSm, 310), font=fSm, fill='#558B2F')
        d.text((cx(1), y + 34 + LH), 'proxy.out → in16' if inp else 'out → proxy.in10', font=fSm, fill='#B71C1C')
        d.text((cx(2), y + 9), 'FIELD', font=fSm, fill='#7B1FA2'); d.text((cx(2), y + 9 + LH), '/ SIM' if inp else '/ P2P', font=fSm, fill='#7B1FA2')
        d.text((cx(3), y + 9), fit(d, R['Point'], fName, 310), font=fName, fill='#5D4037')
        d.text((cx(3), y + 34), fit(d, R['Point_Path'].rsplit('/', 1)[0], fSm, 310), font=fSm, fill='#8D6E63')
        d.text((cx(3), y + 34 + LH), fit(d, r['desc'], F(WS.SAR, 13), 310), font=F(WS.SAR, 13), fill='#6D4C41')
        mid = y + h // 2
        for i in range(4):
            x1, x2 = xs[i] + COLS[i][1], xs[i + 1]
            arrow(d, x1 + 2, x2 - 2, mid, ac, inp)
        bx = xs[4]
        if not r['main']:
            d.text((bx + 4, y + 9), '— แสดงผล / history เท่านั้น (ไม่มี logic block ใช้)', font=fTh, fill='#9E9E9E')
        for j, b in enumerate(r['main'][:4]):
            yy = y + 6 + j * LH; col = ACOL[b['area']]
            d.rectangle((bx, yy + 3, bx + 8, yy + 17), fill=col)
            s1 = f"{b['name']}"; d.text((bx + 14, yy), fit(d, s1, F(DJCB, 14), 300), font=F(DJCB, 14), fill=col)
            w1 = min(d.textlength(s1, font=F(DJCB, 14)), 300) + 22
            s2 = f"{b['type']} · {b['slot']} · {b['folder'].replace('/ChillerPlant/', '').replace('/AirSide/', 'AirSide/')}"
            d.text((bx + w1, yy + 1), fit(d, s2, fSm, 690 - w1), font=fSm, fill='#546E7A')
        if len(r['main']) > 4:
            c = collections.Counter(b['area'] for b in r['main'][4:])
            d.text((bx + 14, y + 6 + 4 * LH), '+ ' + ' · '.join(f'{v} {k}' for k, v in c.items()) + '  (ดูทั้งหมดใน Excel)', font=F(WS.SAR, 14), fill='#78909C')
        y += h + 10
    # legend
    lx = PAD; y += 8
    for k, v in ACOL.items():
        d.rectangle((lx, y + 4, lx + 14, y + 18), fill=v); d.text((lx + 20, y), k, font=fSm, fill='#37474F'); lx += 40 + d.textlength(k, font=fSm)
    arrow(d, lx + 20, lx + 80, y + 11, '#1565C0', True); d.text((lx + 88, y), 'INPUT field → logic', font=fSm, fill='#1565C0')
    arrow(d, lx + 260, lx + 320, y + 11, '#E65100', False); d.text((lx + 328, y), 'OUTPUT logic → field', font=fSm, fill='#E65100')
    im = im.crop((0, 0, WIDTH, y + 30)); im = im.quantize(64); im.save(path, optimize=True)


by = collections.OrderedDict()
for x in rows: by.setdefault(x['r']['Group'], []).append(x)
imgs = []
for g, rs in by.items():
    rs.sort(key=lambda x: (x['r']['Direction'] != 'in', x['r']['Type'], x['r']['Point']))
    parts = [rs[i:i + 10] for i in range(0, len(rs), 10)]
    for k, p in enumerate(parts):
        fn = os.path.join(OUT, 'img', f'IF_{g}_{k + 1}.png')
        R0 = p[0]['r']
        draw_rows(f'{g}  ·  Field interface → Logic' + (f'  ({k + 1}/{len(parts)})' if len(parts) > 1 else ''),
                  f"{R0['Protocol_Name']}  ·  {len(rs)} points  ·  หน้างานต่อ link เฉพาะช่อง 2 (proxy กับ IO_) จุดเดียว  ส่วนช่อง 3-5 ต่อไว้ในโปรแกรมแล้ว", p, fn)
        imgs.append((g, fn))
json.dump(dict(imgs=imgs, groups={g: len(v) for g, v in by.items()},
               counts=collections.Counter(b['area'] for x in rows for b in x['main']),
               nologic=[x['r']['Point'] for x in rows if not x['main']]), open(os.path.join(OUT, 'map.json'), 'w'), ensure_ascii=False, indent=1)
print('xlsx', xl, 'images', len(imgs), 'no-logic points', sum(1 for x in rows if not x['main']))

# example (principle) picture: one input + one output
ex = sorted([x for x in rows if x['r']['Point'] in ('CH01_CdwLeavingTemp', 'CH01StartCmd')], key=lambda x: x['r']['Direction'] != 'in')
if len(ex) == 2:
    draw_rows('หลักการ: 1 point = 1 แถว  (ตัวอย่าง CH01)', 'แถวบน INPUT ลูกศรไปขวา (field ไป logic)  ·  แถวล่าง OUTPUT ลูกศรไปซ้าย (logic ไป field)', ex,
              os.path.join(OUT, 'img', 'IF_00_example.png'))
