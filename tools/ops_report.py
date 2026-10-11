"""ALT_CP_Ref_V7_7 operation & maintenance report (monthly / on demand) -> HTML + XLSX.
Asset health, PM status and forecast, predictive advisories with recommended checks, energy KPI, TOU demand.
usage: ops_report.py --bog config.bog --out <dir> [--hist <dir of History CSV exports>] [--sim v77_fpt.bog] [--site 'name']
  --hist : Workbench History export (CSV) of Smart_*, PM_*, FDD_*, DL_*, Plant_* points; the newest record of each is used
  --sim  : run the offline model for 30 min (demonstration values, not a site record)"""
import argparse, csv, glob, html, os, sys, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from cx_report import bog_values
import zipfile, xml.etree.ElementTree as ET

UNITS = ['CH01', 'CH02', 'CHP01', 'CHP02', 'CDP01', 'CDP02', 'CT01', 'CT02', 'AHU01', 'AHU02']
ADV = {'CH': ('FDD_{u}_EffDegradeAlarm', 'kW/RT above baseline: clean condenser / evaporator tubes, check refrigerant charge, CDW flow, non-condensables'),
       'CHP': ('FDD_{u}_PumpPerfAlarm', 'pump power vs affinity law: check impeller / bearing wear, strainer, valve position, air in system'),
       'CDP': ('FDD_{u}_PumpPerfAlarm', 'pump power vs affinity law: check impeller / bearing wear, strainer, valve position'),
       'CT': ('FDD_{u}_ApproachAlarm', 'tower approach high at full fan: clean fill and nozzles, check fan belt / pitch, basin level'),
       'AHU': (None, '')}
POINTS = [f'{p}_{u}_{s}' for u in UNITS for p, s in (('Smart', 'Health'), ('PM', 'HoursSincePM'), ('PM', 'UsedPct'), ('PM', 'HoursLeft'),
                                                       ('PM', 'DueAlarm'), ('PM', 'OverdueAlarm'), ('PM', 'StartsToday'), ('PM', 'ShortCycleAlarm'),
                                                       ('Smart', 'HealthLowAlarm'))]
POINTS += [ADV[k][0].format(u=u) for u in UNITS for k in ADV if u.startswith(k) and ADV[k][0] and not (k == 'CH' and u.startswith(('CHP',)))]
POINTS += ['FDD_CH01_EffDevPct', 'FDD_CH02_EffDevPct', 'FDD_CHWR_FrozenAlarm', 'FDD_CDR_FrozenAlarm', 'FDD_OA_Temperature_FrozenAlarm',
           'FDD_CHWFlow_FrozenAlarm', 'Smart_MaintAdvisory', 'DL_ActiveSts', 'DL_ChwstOffset', 'Plant_kW', 'Plant_kWh_Today', 'Plant_kW_per_RT_Today',
           'Plant_kWh_Month', 'Plant_kW_per_RT_Month', 'Cooling_RTh_Today']


def adv_of(u):
    for k in ('CHP', 'CDP', 'CH', 'CT', 'AHU'):
        if u.startswith(k): return ADV[k]
    return (None, '')


def hist_values(d):
    out = {}
    for fn in glob.glob(os.path.join(d, '*.csv')):
        rows = list(csv.reader(open(fn, encoding='utf-8-sig', errors='replace')))
        if len(rows) < 2: continue
        head = [h.lower() for h in rows[0]]; vi = next((i for i, h in enumerate(head) if h.startswith('value')), 1)
        v = rows[-1][vi]
        try: v = float(v)
        except ValueError: v = v.lower() in ('true', 'pm due', 'pm overdue', 'short cycling', 'health low')
        out[os.path.splitext(os.path.basename(fn))[0].split('.')[0]] = v
    return out


def sim_values(bog):
    from sim7 import Model6
    from simstart import faithful
    m = faithful(Model6(bog))
    m.run(1800.0, [(0.0, {'PlantEnable': True, 'PlantLoadPct': 70.0}), (1.0, {'BootQualified': True})])
    return {k: m.out.get(k) for k in POINTS if k in m.out}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--bog', required=True); ap.add_argument('--out', required=True); ap.add_argument('--hist'); ap.add_argument('--sim')
    ap.add_argument('--site', default='ALT_CP_Ref_V7_7')
    a = ap.parse_args(); os.makedirs(a.out, exist_ok=True)
    bv = bog_values(a.bog)
    root = ET.fromstring(zipfile.ZipFile(a.bog).read('file.xml'))
    params = {}
    for e in root.iter('p'):
        n = e.get('n') or ''
        if n.startswith(('PM_', 'FDD_', 'DL_', 'SM_')) and e.get('t') in ('c:NumericWritable', 'c:BooleanWritable'):
            fb = next((c for c in e if c.get('n') == 'fallback'), None); v = next((c for c in fb if c.get('n') == 'value'), None) if fb is not None else None
            if v is not None: params[n] = v.get('v')
    hv, src = {}, 'no history export (live values not available)'
    if a.hist: hv.update(hist_values(a.hist)); src = f'history export: {a.hist}'
    if a.sim: hv.update(sim_values(a.sim)); src = 'OFFLINE MODEL 30 min (demonstration only, not a site record)'
    g = lambda k, d=None: hv.get(k, d)
    eq = []
    for u in UNITS:
        an, act = adv_of(u); advon = bool(g(an.format(u=u))) if an else False
        hrs = g(f'PM_{u}_HoursSincePM', 0.0) or 0.0; left = g(f'PM_{u}_HoursLeft', float(params.get(f'PM_{u}_IntervalH', 0)))
        status = 'OVERDUE' if g(f'PM_{u}_OverdueAlarm') else 'DUE' if g(f'PM_{u}_DueAlarm') else 'OK'
        actions = []
        if status != 'OK': actions.append(f'schedule PM ({params.get(f"PM_{u}_IntervalH", "?")} h service)')
        if advon: actions.append(act)
        if g(f'PM_{u}_ShortCycleAlarm'): actions.append('short cycling: review staging deadband / min-run / load')
        eq.append(dict(Equipment=u, Health=f'{(g(f"Smart_{u}_Health") or 0):.0f}', PM_Status=status, Hours_since_PM=f'{hrs:.1f}',
                       PM_Interval_h=params.get(f'PM_{u}_IntervalH', ''), Hours_left=f'{left:.0f}' if isinstance(left, float) else left,
                       Starts_today=f'{(g(f"PM_{u}_StartsToday") or 0):.0f}', Predictive='ADVISORY' if advon else 'OK',
                       Recommended_action='; '.join(actions) or '-'))
    sensors = [(s, 'FROZEN' if g(f'FDD_{s}_FrozenAlarm') else 'OK') for s in ('CHWR', 'CDR', 'OA_Temperature', 'CHWFlow')]
    kpi = [('Plant kW now', g('Plant_kW')), ('Plant kWh today', g('Plant_kWh_Today')), ('kW/RT today', g('Plant_kW_per_RT_Today')),
           ('Plant kWh month', g('Plant_kWh_Month')), ('kW/RT month', g('Plant_kW_per_RT_Month')), ('Cooling RTh today', g('Cooling_RTh_Today')),
           ('CH-01 efficiency deviation %', g('FDD_CH01_EffDevPct')), ('CH-02 efficiency deviation %', g('FDD_CH02_EffDevPct')),
           ('Demand limit', 'enabled' if params.get('DL_Enable') == 'true' else 'disabled'), ('Demand target kW', params.get('DL_TargetKW')),
           ('Demand limiting active now', 'yes' if g('DL_ActiveSts') else 'no'), ('Maintenance advisory', 'YES' if g('Smart_MaintAdvisory') else 'none')]
    fmt = lambda v: f'{v:,.2f}' if isinstance(v, float) else ('-' if v is None else str(v))
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
    n_adv = sum(e['Predictive'] == 'ADVISORY' for e in eq); n_pm = sum(e['PM_Status'] != 'OK' for e in eq)
    lead = f'{n_pm} equipment need PM, {n_adv} predictive advisories, {sum(s[1] == "FROZEN" for s in sensors)} frozen sensors.'
    # XLSX
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Summary'
    ws.append(['O&M report', a.site]); ws.append(['Generated', now]); ws.append(['Source', src]); ws.append(['Result', lead]); ws.append([])
    for k, v in kpi: ws.append([k, fmt(v)])
    ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 60
    s2 = wb.create_sheet('Asset health & PM'); cols = list(eq[0].keys()); s2.append(cols)
    for c in s2[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79')
    for e in eq: s2.append([e[c] for c in cols])
    for i, w in enumerate((10, 8, 10, 14, 13, 10, 12, 12, 80), 1): s2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    s3 = wb.create_sheet('Sensors'); s3.append(['Sensor', 'Status']); [s3.append(list(x)) for x in sensors]
    s4 = wb.create_sheet('Settings'); s4.append(['Parameter', 'Value']); [s4.append([k, v]) for k, v in sorted(params.items())]
    xl = os.path.join(a.out, 'ALT_V7_7_OM_Report.xlsx'); wb.save(xl)
    # HTML
    cls = lambda v: 'bad' if v in ('OVERDUE', 'FROZEN', 'ADVISORY') else 'warn' if v == 'DUE' else ''
    tr = lambda r, cs: '<tr>' + ''.join(f'<td class="{cls(str(r[c]))}">{html.escape(str(r[c]))}</td>' for c in cs) + '</tr>'
    doc = f"""<!doctype html><html><head><meta charset="utf-8"><title>O&M Report {html.escape(a.site)}</title><style>
body{{font:13px/1.45 system-ui,'Sarabun',sans-serif;margin:24px;color:#1b1f24}}h1{{font-size:20px}}h2{{font-size:15px;border-bottom:2px solid #1f4e79;margin-top:24px}}
table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #c9d1d9;padding:4px 6px;text-align:left;vertical-align:top}}th{{background:#1f4e79;color:#fff}}
td.bad{{background:#ffcdd2;font-weight:600}}td.warn{{background:#fff3cd;font-weight:600}}</style></head><body>
<h1>Operation &amp; maintenance report — {html.escape(a.site)}</h1><p><b>{html.escape(lead)}</b><br>Generated {now} · Source: {html.escape(src)}</p>
<h2>1. Energy and plant KPI</h2><table>{''.join(f'<tr><th>{html.escape(k)}</th><td>{html.escape(fmt(v))}</td></tr>' for k, v in kpi)}</table>
<h2>2. Asset health, PM and recommended actions</h2><table><tr>{''.join(f'<th>{c}</th>' for c in cols)}</tr>{''.join(tr(e, cols) for e in eq)}</table>
<h2>3. Sensor health</h2><table><tr><th>Sensor</th><th>Status</th></tr>{''.join(f'<tr><td>{s}</td><td class="{cls(v)}">{v}</td></tr>' for s, v in sensors)}</table>
<h2>4. Sign-off</h2><table><tr><th>Role</th><th>Name</th><th>Signature</th><th>Date</th></tr><tr><td>O&amp;M engineer</td><td></td><td></td><td></td></tr>
<tr><td>Facility manager</td><td></td><td></td><td></td></tr></table></body></html>"""
    hp = os.path.join(a.out, 'ALT_V7_7_OM_Report.html'); open(hp, 'w', encoding='utf-8').write(doc)
    print('report', hp, xl, lead)


if __name__ == '__main__':
    main()
