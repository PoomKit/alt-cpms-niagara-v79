# -*- coding: utf-8 -*-
"""V7.9 equipment register = the single list that says which units / sensors exist on a site.
   ALT_V79_Equipment_Register.xlsx  (sheets: Equipment, Sensors, AirSide, Systems, README)
   * the builder (make_v79.py) creates slots 01..06 for every equipment type and one slot per sensor row;
     'Installed' sets the default of <slot>_Installed in the station (can be changed live in Workbench)
   * Protocol / Device / Address go to the Field IO register (FieldIO/<slot>/IO_<slot>_Protocol + address columns)
   * the same file drives the 3D page (which units are drawn, point mapping) and the 2D / BIM generators"""
import os, json
MAX_UNITS = 6
TYPES = {'CH': ('Chiller', 'chiller'), 'CHP': ('Chilled water pump', 'pump'), 'CDP': ('Condenser water pump', 'pump'), 'CT': ('Cooling tower', 'ct')}
SENSOR_KINDS = {  # kind: unit, valid lo/hi, alarm lo/hi, sim value
    'Temp': ('C', -10.0, 60.0, 5.0, 35.0, 24.0), 'RH': ('RH', 0.0, 100.0, 30.0, 75.0, 55.0), 'Pressure': ('kPa', -50.0, 1600.0, 50.0, 900.0, 350.0),
    'DP': ('kPa', -10.0, 500.0, 20.0, 250.0, 100.0), 'Flow': ('m3h', -10.0, 2000.0, 0.0, 1500.0, 300.0), 'Power': ('kW', -10.0, 2000.0, 0.0, 1500.0, 50.0),
    'Level': ('pct', -5.0, 105.0, 20.0, 90.0, 60.0), 'Generic': ('none', -1e6, 1e6, -1e6, 1e6, 0.0)}
DEFAULT_PROTO = {'CH': 1, 'CHP': 4, 'CDP': 4, 'CT': 4, 'SNS': 5}
RATED = {'CH': '334 TR', 'CHP': '800 USgpm / 22 kW', 'CDP': '1,000 USgpm / 30 kW', 'CT': '420 TR (1 cell, 5.5 kW fan)'}


def default():
    eq = []
    for t, (desc, _) in TYPES.items():
        for i in range(1, MAX_UNITS + 1):
            s = f'{t}{i:02d}'
            eq.append(dict(Slot=s, Type=t, Tag=f'{t}-{i:02d}', Installed='Y' if i <= 2 else 'N', Description=f'{desc} {i}', Rated=RATED[t],
                           Protocol=DEFAULT_PROTO[t], Device='', Address='', Notes='base unit (V7.8 logic)' if i <= 2 else 'expansion slot'))
    sns = []
    for i, (kind, tag, desc, loc) in enumerate([('Temp', 'TT-CHWS-B', 'CHW supply temp, building branch B', 'Plant room'),
                                                ('Pressure', 'PT-CHW-RISER', 'CHW riser pressure, top floor', 'Roof'),
                                                ('Temp', 'TT-PLANTRM', 'Plant room air temperature', 'Plant room'),
                                                ('RH', 'MT-PLANTRM', 'Plant room humidity', 'Plant room'),
                                                ('Flow', 'FT-MU', 'Make-up water flow', 'Plant room'),
                                                ('Power', 'KW-AUX', 'Auxiliary board power', 'Plant room')], 1):
        u, vlo, vhi, alo, ahi, sim = SENSOR_KINDS[kind]
        sns.append(dict(Slot=f'SNS{i:02d}', Installed='N', Tag=tag, Description=desc, Kind=kind, Unit=u, ValidLo=vlo, ValidHi=vhi,
                        AlarmLo=alo, AlarmHi=ahi, SimValue=sim, Location=loc, Protocol=5, Address=''))
    air = [dict(Unit=f'AHU_F{f:02d}', Floor=f, Installed='Y', Zones=4, Description=f'Floor {f} air handling unit') for f in range(2, 9)]
    sysr = [dict(System='fire_alarm', Tag='FACP-01', Type='panel', Floor=1, X=3000, Y=16000, Z=1200, Point='FireAlarmSignal',
                 Note='Fire alarm control panel general alarm (already in the station)')]
    return {'Equipment': eq, 'Sensors': sns, 'AirSide': air, 'Systems': sysr}


COLS = {'AirSide': ['Unit', 'Floor', 'Installed', 'Zones', 'Description'],
        'Systems': ['System', 'Tag', 'Type', 'Floor', 'X', 'Y', 'Z', 'Point', 'Note'], 'Equipment': ['Slot', 'Type', 'Tag', 'Installed', 'Description', 'Rated', 'Protocol', 'Device', 'Address', 'Notes'],
        'Sensors': ['Slot', 'Installed', 'Tag', 'Description', 'Kind', 'Unit', 'ValidLo', 'ValidHi', 'AlarmLo', 'AlarmHi', 'SimValue', 'Location', 'Protocol', 'Address']}


def load(path=None):
    if not path or not os.path.exists(path): return default()
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True); out = {}
    for sh in ('Equipment', 'Sensors', 'AirSide', 'Systems'):
        if sh not in wb.sheetnames: out[sh] = default()[sh]; continue
        ws = wb[sh]; rows = list(ws.iter_rows(values_only=True)); hdr = [str(h).strip() if h else '' for h in rows[0]]
        out[sh] = [dict(zip(hdr, r)) for r in rows[1:] if r and r[0]]
    for r in out['Equipment']:
        r['Type'] = r['Slot'].rstrip('0123456789')
    eq = {r['Slot']: r for r in out['Equipment']}
    for t in TYPES:                                      # slots 01..06 always exist (max equipment); missing rows = not installed
        for i in range(1, MAX_UNITS + 1):
            s = f'{t}{i:02d}'
            if s not in eq: out['Equipment'].append(dict(Slot=s, Type=t, Tag=f'{t}-{i:02d}', Installed='N', Protocol=DEFAULT_PROTO[t]))
    if len(out['Sensors']) < MAX_UNITS:                  # at least 6 spare sensor slots in the station
        have = {r['Slot'] for r in out['Sensors']}
        for i in range(1, MAX_UNITS + 1):
            s = f'SNS{i:02d}'
            if s not in have:
                u, vlo, vhi, alo, ahi, sim = SENSOR_KINDS['Generic']
                out['Sensors'].append(dict(Slot=s, Installed='N', Tag=s, Description='spare sensor slot', Kind='Generic', Unit=u, ValidLo=vlo,
                                           ValidHi=vhi, AlarmLo=alo, AlarmHi=ahi, SimValue=sim, Protocol=5))
    return out


def installed(r):
    return str(r.get('Installed', 'N')).strip().upper() in ('Y', 'YES', 'TRUE', '1')


README = [
    ('ทะเบียนอุปกรณ์ ALT_CP_Ref_V7_9', ''),
    ('ใช้ทำอะไร', 'ไฟล์เดียวที่บอกว่าไซต์นี้มีอุปกรณ์อะไรบ้าง builder อ่านไฟล์นี้แล้วสร้าง station, Field IO, point mapping, 3D และแบบ 2D ให้ตรงกัน'),
    ('Equipment', 'chiller / CHP / CDP / CT อย่างละ 6 slot (01-02 = เครื่องเดิมของ V7.8, 03-06 = slot ขยาย) · Installed = Y/N'),
    ('Sensors', 'sensor เพิ่มเติม (SNS01...) ใส่กี่แถวก็ได้ อย่างน้อย 6 slot · Kind: Temp, RH, Pressure, DP, Flow, Power, Level, Generic'),
    ('AirSide', 'AHU ประจำชั้น 2-8 (ชั้น 1 = AHU01/02, OAU01, FCU01-04 ของ V7.8) · Installed = Y/N'),
    ('Systems', 'ระบบในอนาคต (fire alarm, fire protection, lighting, power) · ใส่อุปกรณ์ทีละแถว: ชั้น, ตำแหน่ง X/Y/Z (มม.), Point ใน station → ขึ้นใน 3D เป็นจุดพร้อมสถานะ'),
    ('เพิ่มอุปกรณ์ (วิธีที่ 1, ไม่ต้องสร้างใหม่)', 'ใน Workbench ตั้ง <slot>_Installed = true (เช่น CH03_Installed) → logic นับเครื่องนั้นทันที → ต่อ Field IO กลุ่ม <slot> → ทดสอบ P2P'),
    ('เพิ่มอุปกรณ์ (วิธีที่ 2, สร้างใหม่)', 'แก้ไฟล์นี้ (Installed = Y, Protocol, Device) → สั่ง make_v79.py --register ไฟล์นี้ → ได้ config.bog + ทะเบียนจุด + 3D ใหม่ที่ตรงกัน'),
    ('ลดอุปกรณ์', 'ตั้ง Installed = N (หรือ <slot>_Installed = false ใน Workbench) → เครื่องนั้นไม่ถูกเลือก ไม่ถูกสั่งเดิน และเป็นเงาจางใน 3D'),
    ('Protocol', '1 BACnet/IP · 2 BACnet MS/TP · 3 Modbus TCP · 4 Modbus RTU · 5 hard-wired I/O'),
]


def write_template(path, data=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    data = data or default(); wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
    hf, fill, thin = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1B2A3A'), Side(style='thin', color='C9D0D8')
    for i, (a, b) in enumerate(README, 1):
        ws.cell(i, 1, a).font = Font(bold=True, size=14 if i == 1 else 11); ws.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 110
    for sh in ('Equipment', 'Sensors', 'AirSide', 'Systems'):
        w = wb.create_sheet(sh); cols = COLS[sh]
        for j, c in enumerate(cols, 1):
            x = w.cell(1, j, c); x.font, x.fill = hf, fill; w.column_dimensions[x.column_letter].width = max(10, min(46, len(c) + 6))
        for i, r in enumerate(data[sh], 2):
            for j, c in enumerate(cols, 1):
                x = w.cell(i, j, r.get(c)); x.border = Border(bottom=thin)
        w.freeze_panes = 'A2'; w.auto_filter.ref = w.dimensions
        if 'Installed' in cols:
            dv = DataValidation(type='list', formula1='"Y,N"', allow_blank=False); w.add_data_validation(dv)
            k = cols.index('Installed') + 1; L = w.cell(1, k).column_letter; dv.add(f'{L}2:{L}500')
        if 'Protocol' in cols:
            dv = DataValidation(type='whole', operator='between', formula1='1', formula2='5'); w.add_data_validation(dv)
            k = cols.index('Protocol') + 1; L = w.cell(1, k).column_letter; dv.add(f'{L}2:{L}500')
    wb.save(path)
