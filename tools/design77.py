# -*- coding: utf-8 -*-
"""Site deployment design (Thai) for ALT_CP_Ref_V7_7 -> DOCX.  usage: design77.py register.csv arch.png out.docx"""
import sys, os, csv, collections, math
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REG, ARCH, OUT = sys.argv[1:4]
FONT = os.environ.get('MANUAL_FONT', 'Leelawadee UI'); SC = float(os.environ.get('MANUAL_SCALE', '1'))
doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(1.8))
W = 21.0 - 3.6


def style(st, size, bold=None, color=None):
    st.font.name = FONT; st.font.size = Pt(size * SC)
    if bold is not None: st.font.bold = bold
    if color: st.font.color.rgb = RGBColor.from_string(color)
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(k), FONT)
    e = OxmlElement('w:szCs'); e.set(qn('w:val'), str(int(size * SC * 2))); rpr.append(e)
    if bold: rpr.append(OxmlElement('w:bCs'))


style(doc.styles['Normal'], 10)
for l, s_, c in ((1, 16, '1F4E79'), (2, 13, '1F4E79'), (3, 11, '2E5E8C')): style(doc.styles[f'Heading {l}'], s_, True, c)


def run(p, text, size=None, bold=False, color=None):
    r = p.add_run(text); r.bold = bold
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size * SC); rpr = r._r.get_or_add_rPr(); e = OxmlElement('w:szCs'); e.set(qn('w:val'), str(int(size * SC * 2))); rpr.append(e)
    return r


def P(text, size=None, bold=False, color=None, align=None):
    p = doc.add_paragraph()
    for i, part in enumerate(text.split('**')): run(p, part, size, bold or i % 2 == 1, color)
    if align: p.alignment = align
    return p


def BL(items):
    for t in items:
        p = doc.add_paragraph(style='List Bullet')
        for i, part in enumerate(t.split('**')): run(p, part, None, i % 2 == 1)


def NL(items):
    for t in items:
        p = doc.add_paragraph(style='List Number')
        for i, part in enumerate(t.split('**')): run(p, part, None, i % 2 == 1)


def shade(c, h):
    tcPr = c._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), h); tcPr.append(s)


def T(head, rows, widths, size=8.5):
    t = doc.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for c, h, w in zip(t.rows[0].cells, head, widths):
        c.width = Cm(w); run(c.paragraphs[0], h, size, True, 'FFFFFF'); shade(c, '1F4E79')
    for k, r in enumerate(rows):
        for c, v, w in zip(t.add_row().cells, r, widths):
            c.width = Cm(w); run(c.paragraphs[0], str(v), size)
            if k % 2: shade(c, 'F2F5F8')
    doc.add_paragraph()


# ---------------------------------------------------------------- data
reg = list(csv.DictReader(open(REG, encoding='utf-8-sig')))
cnt = collections.Counter((r['Protocol_Name'], r['Type']) for r in reg)
grp = collections.defaultdict(set)
for r in reg: grp[r['Protocol_Name']].add(r['Group'])
hw = collections.Counter(r['Type'] for r in reg if r['Protocol'] == '5')
spare = lambda n: math.ceil(n * 1.2)

# ---------------------------------------------------------------- content
P('ALT_CP_Ref_V7_7', 24, True, '1F4E79', WD_ALIGN_PARAGRAPH.CENTER)
P('แบบระบบอย่างง่ายสำหรับนำ logic ไปใช้กับอุปกรณ์จริงหน้างาน', 15, True, align=WD_ALIGN_PARAGRAPH.CENTER)
P('Site deployment design · chiller plant + air side · Niagara N4 · r5 (2 ต.ค. 2026)', 10, align=WD_ALIGN_PARAGRAPH.CENTER)
P('เอกสารอ้างอิงสำหรับออกแบบโครงการ ต้องยืนยันรุ่นอุปกรณ์ จำนวนจุด และ address กับผู้ผลิตก่อนสั่งซื้อ', 9, color='B7791F', align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_heading('1. สรุป', 1)
P(f'logic ทั้งหมดรันบน plant controller ตัวเดียว (Niagara JACE-9000 หรือ Niagara Edge) ต่ออุปกรณ์ด้วย 4 ช่องทาง: BACnet/IP (chiller, DDC panel), '
  f'Modbus RTU (VFD, มิเตอร์), BACnet MS/TP (FCU, zone) และ IO module (sensor และ interlock แบบ hard-wired) station มีจุดเชื่อมพร้อมแล้ว {len(reg)} จุด '
  'ใน folder FieldIO แค่ต่อ driver point แล้วสลับเป็น FIELD ทีละกลุ่ม')
doc.add_picture(ARCH, width=Cm(W)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
P('รูป 1 สถาปัตยกรรมระบบ', 8, color='5F6B7A', align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_heading('2. จำนวนจุดและขนาดอุปกรณ์ควบคุม', 1)
T(['Protocol', 'AI', 'DI', 'DO', 'AO', 'รวม', 'กลุ่มอุปกรณ์'],
  [(pn, cnt[(pn, 'AI')], cnt[(pn, 'DI')], cnt[(pn, 'DO')], cnt[(pn, 'AO')], sum(cnt[(pn, t)] for t in ('AI', 'DI', 'DO', 'AO')), ', '.join(sorted(grp[pn])))
   for pn in ('BACnet/IP', 'BACnet MS/TP', 'Modbus RTU', 'Hard-wired I/O')], [2.6, 1.0, 1.0, 1.0, 1.0, 1.2, 9.6], 8)
P(f'**IO module (hard-wired) เผื่อ 20 %:** universal input {spare(hw["AI"] + hw["DI"])} จุด (AI {hw["AI"]} + DI {hw["DI"]}), DO {spare(hw["DO"])} จุด, AO {spare(hw["AO"])} จุด '
  '→ ประมาณ 12 × module 16 UI, 2 × module 8 DO, 3 × module 8 AO แบ่งตามตู้ (ห้องเครื่อง, AHU, OAU, พัดลม)')
P('**ข้อแนะนำเพื่อลดสายไฟ:** ข้อมูล VFD (Hz, A, kW, fault) อ่านผ่าน Modbus จาก VFD โดยตรง ส่วน hard-wire เฉพาะคำสั่ง start, run status, '
  'transmitter ของกระบวนการ (TT, PT, flow, level) และสัญญาณ safety ค่า default ใน V7_7_FIELD_IO_REGISTER.csv เป็นจุดเริ่ม แก้ protocol ราย point ได้')
T(['รายการ', 'สเปกแนะนำ', 'จำนวน'], [
    ('Plant controller', 'Niagara JACE-9000 (หรือ Edge) · Niagara 4.15 · license ≥ 1,000 points / 50 devices', '1'),
    ('RS-485 port', 'ใช้พอร์ตบน JACE + option card RS-485', '3 bus'),
    ('IO module', 'IO-R / IO-16 series หรือเทียบเท่าที่ Niagara รองรับ', 'ตามข้อ 2'),
    ('Network switch', 'Managed switch 8–16 port, VLAN, DIN rail', '1–2'),
    ('UPS', '1 kVA online, ≥ 30 นาที สำหรับ JACE + switch + IO', '1'),
    ('Workstation', 'Workbench + browser, หรือ Supervisor (หลายอาคาร)', '1')], [3.8, 11.0, 2.6], 8.5)

doc.add_heading('3. การสื่อสารและ network', 1)
T(['ช่องทาง', 'อุปกรณ์', 'ค่าตั้ง', 'หมายเหตุ'], [
    ('Ethernet VLAN BMS 10.10.10.0/24', 'JACE .10 · CH-01 .21 · CH-02 .22 · AHU/OAU DDC .31–.33 · Supervisor .5 · AI server .6', 'BACnet/IP UDP 47808, Fox TLS 4911, HTTPS 443',
     'แยก VLAN จาก IT, ปิดพอร์ตที่ไม่ใช้'),
    ('RS-485 #1 Modbus RTU', 'VFD CHP-01/02, CDP-01/02, CT fan ×4 · ID 1–8', '19,200 bps 8N1', 'สาย twisted pair มี shield, terminator 120 Ω ปลายสาย'),
    ('RS-485 #2 Modbus RTU', 'power meter, BTU meter, make-up meter · ID 11–20', '9,600 bps 8N1', 'ระยะรวม ≤ 1,000 m, ≤ 32 node'),
    ('RS-485 #3 BACnet MS/TP', 'FCU-01…04 MAC 1–4, zone gateway MAC 10', '38,400 bps, max master 15', 'ตั้ง device instance ไม่ซ้ำ'),
    ('IO bus', 'IO module ในตู้ห้องเครื่อง / AHU', 'ตามผู้ผลิต', 'TT 10k NTC / PT1000, PT 4–20 mA, DO ผ่าน interposing relay')], [3.6, 6.0, 3.6, 4.2], 8)

doc.add_heading('4. อุปกรณ์ field ที่ต้องมี', 1)
T(['อุปกรณ์', 'สเปกแนะนำ', 'ใช้กับ logic'], [
    ('Temperature (TT)', 'immersion + thermowell, PT1000 / 10k NTC, ±0.2 °C', 'CHWS/R, CDS/R, chiller in/out, AHU SAT/RAT'),
    ('Pressure (PT) / DP', '4–20 mA 0–10 bar (pump), DP 0–2.5 bar ปลาย loop', 'pump HP trip, DP control (W-01)'),
    ('Flow meter', 'electromagnetic, Modbus หรือ 4–20 mA', 'min-flow bypass, staging, BTU'),
    ('Flow switch', 'paddle / thermal, contact ต่อทั้ง chiller panel และ BMS', 'C-04 flow-switch trip'),
    ('Power / BTU meter', 'Modbus RTU, class 0.5S (power), class 2 (BTU)', 'kW/RT, FDD efficiency, demand limit'),
    ('Refrigerant detector', '4–20 mA + relay 2 ระดับ (300 / 1,000 ppm)', 'S-01 refrigerant leak'),
    ('Basin level, conductivity', 'ultrasonic 4–20 mA + float LL/HH, conductivity 4–20 mA', 'T-03, T-04 water treatment'),
    ('OA temp / RH', '4–20 mA, shield จากแดด', 'wet bulb, CDS reset, CO2 / dehum')], [3.6, 7.6, 6.2], 8)
P('**Safety ต้อง hard-wire ไม่พึ่ง BMS:** E-stop chain, flow switch ไปที่ chiller control panel และ fire alarm ไปที่ตู้พัดลม BMS เป็นผู้เฝ้าดูและ interlock ชั้นที่สอง')

doc.add_heading('5. ฟังก์ชันอัจฉริยะ (r5) และการตั้งค่าหน้างาน', 1)
T(['ฟังก์ชัน', 'ทำอะไร', 'ตั้งค่าหน้างาน'], [
    ('PM ตามชั่วโมงเดิน', 'นับชั่วโมงตั้งแต่ PM ครั้งล่าสุด เตือนที่ 90 % และ 100 % ของรอบ, ปุ่ม PM DONE', 'PM_<อุปกรณ์>_IntervalH ตามแผน PM ของผู้ผลิต'),
    ('Starts / short cycling', 'นับการ start ต่อวัน เตือนเมื่อเกิน', 'SM_ChillerMaxStartsDay (6), SM_EquipMaxStartsDay (12)'),
    ('Chiller efficiency drift', 'kW/RT กรองช้าเทียบ baseline เตือนเมื่อสูงกว่า 15 % นาน 2 ชม.', 'FDD_CHxx_BaselineKwPerRt จากข้อมูล 2 สัปดาห์หลัง commissioning (ท่อสะอาด)'),
    ('Pump hydraulic health', 'kW จริงเทียบ affinity law (kW rated × speed³)', 'FDD_<pump>_BaselineRatio จากข้อมูลช่วงแรก'),
    ('Tower degradation', 'approach สูงกว่า design + 2 K ขณะพัดลม ≥ 90 %', 'CT_Approach, FDD_CtApproachMarginK'),
    ('Frozen sensor', 'ค่าไม่เปลี่ยนเลย 3 ชม. ขณะ plant เดิน', 'FDD_FrozenTol'),
    ('Health score 0–100', 'รวม fault, PM, short cycle, predictive เป็นคะแนนเดียวต่ออุปกรณ์', 'FDD_HealthLowLimit (70)'),
    ('TOU demand limit', 'ช่วง peak จ.–ศ. 09:00–22:00 ถ้า plant kW > เป้า ยก CHWST สูงสุด +2 K', 'DL_Enable, DL_TargetKW จากบิลค่าไฟ 12 เดือน, เพิ่มวันหยุดราชการใน Sch_TouPeak'),
    ('O&M report', 'tools/ops_report.py → HTML / XLSX: health, PM, action, sensor, KPI', 'export history ทุกเดือน')], [3.6, 7.2, 6.6], 8)

doc.add_heading('6. ขั้นตอนนำไปใช้หน้างาน', 1)
NL(['**สำรวจและยืนยันจุด:** ใช้ V7_7_FIELD_IO_REGISTER.csv เทียบกับ points list ของผู้ผลิต chiller / VFD / มิเตอร์ แก้ protocol และ address',
    '**ออกแบบตู้และสาย:** IO module, power 24 VAC, interposing relay, surge protection, RS-485 topology แบบ daisy chain',
    '**ติดตั้ง JACE:** Niagara 4.15, license, certificate TLS, ผู้ใช้แยกบัญชีตามบทบาท, ปิด service ที่ไม่ใช้',
    '**คัดลอก station:** ALT_CP_Ref_V7_7 (config.bog + shared/px) ตั้งชื่อ station ตามโครงการ เปลี่ยน FoxS port ถ้าจำเป็น',
    '**สร้าง driver:** BacnetNetwork / ModbusAsyncNetwork ตามไฟล์จาก fieldio_export.py, discover device, สร้าง proxy point',
    '**link เข้า IO_:** input: proxy.out → IO_<point>.in16 · output: IO_<point>.out → proxy.in10, ต่อสถานะ device ไปที่ IO_<กลุ่ม>_CommFault',
    '**ปลดล็อกและสลับ:** IO_DriversInstalled = true แล้วสลับ Use field ทีละกลุ่ม ตรวจป้ายหัวหน้าเป็น FIELD I/O',
    '**P2P:** TEST MODE + ISOLATE ทีละกลุ่ม ใช้ V7_7_P2P_CHECKLIST.csv บันทึกผล',
    '**Functional test:** RUN ALL FT01–FT10 ในช่วงเวลาที่ตกลงกับเจ้าของอาคาร ออกรายงานด้วย cx_report.py',
    '**ปรับจูน:** PID loops, ค่า trip / alarm ตามผู้ผลิต, ตั้ง baseline FDD หลังเดินจริง 2 สัปดาห์, ตั้ง PM interval และ TOU target',
    '**สำรองและส่งมอบ:** backup station (BackupService), อบรมผู้ใช้ด้วย Operation Manual, O&M report เดือนแรก'])

doc.add_heading('7. ความปลอดภัยไซเบอร์', 1)
BL(['บัญชีผู้ใช้แยกรายคน, รหัสผ่านแข็งแรง, role: operator / engineer / admin', 'ใช้ Fox TLS และ HTTPS เท่านั้น, ปิด Fox / HTTP แบบไม่เข้ารหัส',
    'แยก VLAN BMS, firewall ให้ remote access ผ่าน VPN เท่านั้น', 'backup อัตโนมัติรายสัปดาห์ และเก็บสำเนาไว้นอกเครื่อง', 'อัปเดต Niagara patch ตามรอบของผู้ผลิต'])

doc.add_heading('8. สิ่งที่ต้องยืนยันก่อนนำไปใช้', 1)
BL(['รุ่นและ points list ของ chiller, VFD, มิเตอร์ (address ใน register เป็นค่าตั้งต้น)', 'ค่า trip / alarm / PM interval ตามคู่มือผู้ผลิตและแผน PM ของอาคาร',
    'อัตราค่าไฟและช่วง TOU ของผู้ใช้ไฟ (MEA / PEA) และเป้า demand จากบิลจริง', 'baseline ของ FDD ต้องเก็บจากข้อมูลจริงหลังท่อสะอาด ค่าเริ่มต้นมาจาก simulation',
    'license Niagara (จำนวน point / device) ตามจำนวนจริงหลังสำรวจ'])
doc.add_heading('แหล่งข้อมูล', 2)
BL(['แนวโน้ม predictive maintenance / FDD ของ chiller plant: oxmaint.com (Chiller plant FDD platform, Chiller predictive maintenance AI & IoT), facilityexecutive.com (Fault detection and diagnostics transforms facility operations)',
    'อัตรา TOU ไทย (peak จ.–ศ. 09:00–22:00) และ demand charge จาก peak 15 นาที: capsolar.co.th (TOU schedule, demand charge), greenenergythailand.com (Time-of-use rates in Thailand)'])
doc.save(OUT); print('saved', OUT)
