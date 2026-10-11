# -*- coding: utf-8 -*-
"""AI (cloud) tag handover for ALT_CP_Ref_V7_7: every point carrying melco:pointNameString, with access (READ / WRITE),
Niagara ord, type, unit, range, write slot and what the logic does with it.
usage: ai_tag_map.py <config.bog> <regdir> <outdir>   -> AI_TAG_MAP_V7_7.xlsx / .csv / .json + ALT_CP_V7_7_AI_Cloud_Interface.docx
env MANUAL_FONT / MANUAL_SCALE as the other documents (Sarabun 1.15 for PDF conversion)."""
import sys, os, csv, json, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ws_manual as WS
from v77_cloud import CLOUD_W

BOG, REGDIR, OUT = sys.argv[1:4]; os.makedirs(OUT, exist_ok=True)
comp, fol, links = WS.load(BOG)
rev = collections.defaultdict(list)
for t, lk in links.items():
    for s, src, _ in lk:
        if src: rev[src].append((t, s))
pl = {r['Point']: r for r in csv.DictReader(open(os.path.join(REGDIR, 'V7_POINT_LIST.csv'), encoding='utf-8-sig'))}
tags = [r for r in csv.DictReader(open(os.path.join(REGDIR, 'V7_TAG_REGISTER.csv'), encoding='utf-8-sig')) if r['Tag'] == 'melco:pointNameString']
fio = {r['Point']: r for r in csv.DictReader(open(os.path.join(REGDIR, 'V7_7_FIELD_IO_REGISTER.csv'), encoding='utf-8-sig'))}
fb = lambda n: WS.fallback_of(comp[n]) if n in comp else None
num = lambda n: float(fb(n)) if fb(n) not in (None, '') else None
CLS = {'PV': 'ค่าวัด (process value)', 'STS': 'สถานะ', 'SP': 'ค่าตั้ง (operator / ใช้งานจริง)', 'CSP': 'ค่าสั่ง (command / setpoint)'}
RANGE = {'CHW_LeavingTemperatureSetpoint_AI': ('AI_ChwstMin', 'AI_ChwstMax'), 'CHWDP_Setpoint_AI': ('AI_DpMin', 'AI_DpMax'),
         'CDS_Setpoint_AI': ('CDW_MinSetpoint', 'CDW_MaxSetpoint')}
MEAN = {'HeartbeatTimestampSetpoint': 'AI เขียนค่าที่เปลี่ยนทุกครั้ง (แนะนำ Unix time วินาที) ทุก 60 s (ห้ามห่างเกิน 5 นาที)',
        'CHW_LeavingTemperatureSetpoint_AI': 'CHW supply setpoint ที่ AI ต้องการ', 'CHWDP_Setpoint_AI': 'CHW differential pressure setpoint ที่ AI ต้องการ',
        'CDS_Setpoint_AI': 'condenser water supply setpoint ที่ AI ต้องการ',
        'HeartbeatDetectorStatus': 'true = station เห็น heartbeat ของ AI ปกติ', 'AI_Available': 'true = heartbeat ปกติ และค่า AI ทั้ง 3 อยู่ในช่วง',
        'ModeCmd_UI': 'operator เลือก true = ให้ AI คุม / false = LOCAL (AI อ่านอย่างเดียว)',
        'OperationMode': 'true = plant ใช้ค่าจาก AI อยู่ตอนนี้ / false = LOCAL', 'PlantEnable': 'operator เปิด / ปิด plant',
        'PlantRunRequest': 'plant ต้องเดิน (คำนวณจาก logic)', 'CHW_LeavingTemperatureSetpoint_LOCAL': 'CHWST setpoint ของ LOCAL mode',
        'CHW_LeavingTemperatureSetpoint_ACTIVE': 'CHWST setpoint ที่ใช้คุมจริงตอนนี้ (AI หรือ LOCAL + TOU offset)',
        'CHWDP_Setpoint_ACTIVE': 'DP setpoint ที่ใช้คุมจริง', 'CDS_Setpoint_ACTIVE': 'CDS setpoint ที่ใช้คุมจริง'}
rows = []
for t in tags:
    n = t['Point']; v = t['Value']; el = comp.get(n)
    typ = el.get('t').split(':')[1].replace('Writable', '') if el is not None else '?'
    code = v.rsplit('_', 1)[-1]; w = n in CLOUD_W
    p = pl.get(n, {}); unit = p.get('Unit', '') or ('true/false' if typ == 'Boolean' else '')
    lo = hi = ''
    if n in RANGE: lo, hi = num(RANGE[n][0]), num(RANGE[n][1])
    src = ('AI cloud เขียน' if w else 'field point (' + fio[n]['Protocol_Name'] + ')' if n in fio else
           'operator / ค่าตั้ง' if p.get('Kind') in ('param', 'bparam', 'mode') else 'คำนวณจาก logic')
    if n in fio and fio[n]['Direction'] == 'out': src = 'คำสั่งจาก logic ไปอุปกรณ์ (' + fio[n]['Protocol_Name'] + ')'
    used = [x for x, _ in rev.get(n, []) if not x.startswith(('IOS_', 'Gfx_', 'Sim_', 'Obs_', 'CX_'))][:6]
    note = MEAN.get(n, p.get('Note', ''))
    if code == 'CSP' and not w: note = (note + ' · ' if note else '') + 'AI อ่านอย่างเดียว (station คำนวณ / สั่งเอง ผ่าน interlock)'
    rows.append(dict(No=v.split('_', 1)[0], Tag=v, Access='WRITE' if w else 'READ', Point=n, Ord='station:|slot:' + p.get('Slot_Path', t['Slot_Path']),
                     Type=typ, Unit=unit, Class=f'{code} {CLS.get(code, "")}', Min=lo, Max=hi, Write_slot='in10 (priority 10)' if w else '',
                     Fallback_if_not_written=fb(n) if w else '', Value_source=src, Meaning=note, Used_by_logic=', '.join(used)))
rows.sort(key=lambda r: r['No'])
cols = list(rows[0])
with open(os.path.join(OUT, 'AI_TAG_MAP_V7_7.csv'), 'w', newline='', encoding='utf-8-sig') as f:
    wr = csv.DictWriter(f, fieldnames=cols); wr.writeheader(); wr.writerows(rows)
contract = dict(site='ALT_CP_Ref_V7_7', tag='melco:pointNameString', transport='cloud connector (no BACnet / field wiring for AI)',
                write_slot='in10', heartbeat=dict(tag=next(r['Tag'] for r in rows if r['Point'] == 'HeartbeatTimestampSetpoint'),
                                                  value='must change every write (Unix time seconds recommended)', period_s=60, max_period_s=300,
                                                  pre_disconnect_min=num('HB_Step1_PreDisconnectMin'), disconnect_min=num('HB_Step2_DisconnectMin'),
                                                  beats_needed_after_connect=num('HB_RequiredBeats')),
                ai_in_control_when=['ModeCmd_UI = true (operator)', 'AI_SourceCloud = true (site setting)', 'heartbeat OK', 'all 3 AI setpoints within Min..Max'],
                points=[{k: r[k] for k in ('Tag', 'Access', 'Point', 'Ord', 'Type', 'Unit', 'Min', 'Max', 'Write_slot')} for r in rows])
json.dump(contract, open(os.path.join(OUT, 'AI_TAG_MAP_V7_7.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Tag map'; ws.append(cols)
for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79'); c.alignment = Alignment(wrap_text=True)
for r in rows:
    ws.append([r[c] for c in cols])
    if r['Access'] == 'WRITE':
        for c in ws[ws.max_row]: c.fill = PatternFill('solid', fgColor='FFF3E0'); c.font = Font(bold=True)
for i, wdt in enumerate((5, 34, 8, 36, 62, 9, 9, 26, 6, 6, 16, 10, 30, 60, 50), 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
for row_ in ws.iter_rows(min_row=2):
    for c in row_: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
RULES = [
    ('การเชื่อมต่อ', 'AI เชื่อมจาก cloud ผ่าน cloud connector ของ station ไม่ใช้สาย BACnet / field wiring'),
    ('การหา point', f'ใช้ tag melco:pointNameString (ช่อง Tag) เป็น key อย่าใช้ชื่อ / path ของ Niagara เพราะอาจเปลี่ยน · มีทั้งหมด {len(rows)} tag'),
    ('WRITE', f'AI เขียนได้ {sum(r["Access"] == "WRITE" for r in rows)} point เท่านั้น (แถวสีส้ม) ที่ slot in10 · point อื่นอ่านอย่างเดียว'),
    ('Heartbeat', 'เขียน 021 ทุก 60 s ด้วยค่าที่เปลี่ยนทุกครั้ง (Unix time) · ห่างเกิน 6 นาที = pre-disconnect alarm · เกิน 12 นาที = disconnect → plant กลับ LOCAL'),
    ('เริ่มคุม', 'ต้องมี heartbeat อย่างน้อย 2 ครั้ง และ setpoint ทั้ง 3 (003, 006, 008) อยู่ในช่วง Min..Max พร้อมกัน ถ้าค่าใดค่าหนึ่งหลุดช่วง plant ใช้ LOCAL ทั้งหมด'),
    ('เงื่อนไขฝั่ง station', 'operator ตั้ง ModeCmd_UI (024) = true และหน้างานตั้ง AI_SourceCloud = true · อ่านผลได้จาก 025 OperationMode (true = AI คุมอยู่) และ 023 AI_Available'),
    ('ค่าที่ไม่ได้เขียน', 'ถ้า AI ยังไม่เคยเขียน setpoint point จะใช้ค่า fallback (ช่อง Fallback) ควรเขียนทั้ง 3 ค่าทุกครั้งพร้อม heartbeat'),
    ('ค่าที่ใช้คุมจริง', 'ดู *_ACTIVE (005, 007, 009) · CHWST ACTIVE อาจสูงกว่าค่า AI ได้ถึง 2 K ช่วง TOU peak ถ้าเปิด demand limit'),
    ('คำสั่งอุปกรณ์ (CSP ของ CH / CHP / CDP / CT / valve)', 'AI อ่านอย่างเดียว station เป็นผู้สั่งผ่าน sequence และ interlock ถ้าต้องการให้ AI สั่งเปิด / ปิดเครื่องหรือความเร็วโดยตรง ต้องเพิ่ม AI request point พร้อม guard ก่อน'),
]
s2 = wb.create_sheet('AI rules'); s2.append(['หัวข้อ', 'กติกา']); [s2.append(list(x)) for x in RULES]
s2.column_dimensions['A'].width = 34; s2.column_dimensions['B'].width = 140
for c in s2[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E79')
wb.save(os.path.join(OUT, 'AI_TAG_MAP_V7_7.xlsx'))

# ---------------------------------------------------------------- DOCX (handover to the AI team)
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
FONT = os.environ.get('MANUAL_FONT', 'Leelawadee UI'); SC = float(os.environ.get('MANUAL_SCALE', '1'))
doc = Document(); sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Cm(1.4))
W = 29.7 - 2.8


def style(st, size, bold=None, color=None):
    st.font.name = FONT; st.font.size = Pt(size * SC)
    if bold is not None: st.font.bold = bold
    if color: st.font.color.rgb = RGBColor.from_string(color)
    rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(k), FONT)
    e = OxmlElement('w:szCs'); e.set(qn('w:val'), str(int(size * SC * 2))); rpr.append(e)


style(doc.styles['Normal'], 10)
for l, s_ in ((0, 22), (1, 15), (2, 12)): style(doc.styles['Title' if l == 0 else f'Heading {l}'], s_, True, '1F4E79')


def run(p, text, size=None, bold=False, color=None):
    r = p.add_run(text); r.bold = bold
    if color: r.font.color.rgb = RGBColor.from_string(color)
    if size:
        r.font.size = Pt(size * SC); e = OxmlElement('w:szCs'); e.set(qn('w:val'), str(int(size * SC * 2))); r._r.get_or_add_rPr().append(e)


def shade(c, h):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), h); c._tc.get_or_add_tcPr().append(s)


def T(head, data, widths, size=8, hl=None):
    t = doc.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for c, h, w in zip(t.rows[0].cells, head, widths): c.width = Cm(w); run(c.paragraphs[0], h, size, True, 'FFFFFF'); shade(c, '1F4E79')
    for i, w in enumerate(widths): t.columns[i].width = Cm(w)
    for k, r in enumerate(data):
        for c, v, w in zip(t.add_row().cells, r, widths):
            c.width = Cm(w); run(c.paragraphs[0], str(v), size, bool(hl and hl(r)))
            if hl and hl(r): shade(c, 'FFE0B2')
            elif k % 2: shade(c, 'F2F5F8')
    doc.add_paragraph()


doc.add_heading('AI Cloud Interface – Tag Map', 0)
p = doc.add_paragraph(); run(p, 'ALT_CP_Ref_V7_7 (r7) · สำหรับทีม AI · REFERENCE SIMULATION – NOT PROJECT APPROVED', 10, True, 'B71C1C')
doc.add_heading('1. กติกาการเชื่อมต่อ', 1)
T(['หัวข้อ', 'กติกา'], RULES, (6, W - 6), 9)
doc.add_heading('2. Point ที่ AI เขียน (WRITE)', 1)
wr_ = [r for r in rows if r['Access'] == 'WRITE']
T(['Tag', 'Point', 'Type / unit', 'Min', 'Max', 'Slot', 'ถ้าไม่เขียน', 'ความหมาย'],
  [[r['Tag'], r['Point'], f"{r['Type']} {r['Unit']}", r['Min'], r['Max'], r['Write_slot'], r['Fallback_if_not_written'], r['Meaning']] for r in wr_],
  (5.4, 5.6, 2.4, 1.2, 1.2, 2.6, 1.6, W - 20), 8.5)
doc.add_heading('3. Point ที่ AI อ่าน (READ)', 1)
T(['Tag', 'Point', 'Type / unit', 'ชนิด', 'ที่มา', 'ความหมาย'],
  [[r['Tag'], r['Point'], f"{r['Type']} {r['Unit']}", r['Class'], r['Value_source'], r['Meaning']] for r in rows if r['Access'] == 'READ'],
  (5.4, 6.2, 2.4, 3.8, 4.2, W - 22), 7.5, hl=lambda r: r[1] in ('OperationMode', 'AI_Available', 'HeartbeatDetectorStatus'))
doc.add_heading('4. ลำดับการทำงาน', 1)
for s in ['AI เขียน heartbeat (021) ทุก 60 s ตลอดเวลา และเขียน setpoint 003 / 006 / 008 ทุกรอบที่คำนวณ',
          'station ตรวจ heartbeat (ต้องได้ ≥ 2 ครั้ง) และช่วงค่าของ setpoint ทั้ง 3 → AI_Available (023) = true',
          'ถ้า operator เลือก AI (ModeCmd_UI 024 = true) → OperationMode (025) = true → *_ACTIVE (005 / 007 / 009) ใช้ค่าจาก AI (ไล่ค่าแบบ ramp)',
          'heartbeat หยุด > 6 นาที → alarm pre-disconnect · > 12 นาที → disconnect → plant กลับ LOCAL ทันที ไม่ต้องให้ AI สั่ง',
          'AI ตรวจผลจาก 025 / 023 / 022 และค่า *_ACTIVE ทุกครั้ง อย่าสมมติว่าค่าที่เขียนถูกใช้']:
    q = doc.add_paragraph(style='List Number'); run(q, s)
doc.save(os.path.join(OUT, 'ALT_CP_V7_7_AI_Cloud_Interface.docx'))
print('tags', len(rows), 'write', len(wr_), OUT)
